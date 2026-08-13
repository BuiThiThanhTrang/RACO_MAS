from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\SideProject\ChatDev\output\bao_cao_tien_do_puppeteer_generalization.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="A6A6A6", sz="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), sz)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_borders(table, color="D9E2F3", sz="6"):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, color=color, sz=sz)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def add_paragraph(doc, text="", style=None, bold=False, color=None, size=None, italic=False):
    paragraph = doc.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)
    return paragraph


def add_placeholder(doc, title, note):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F7")
    set_cell_border(cell, color="7F7F7F", sz="10")
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(31, 77, 120)

    note_paragraph = cell.add_paragraph()
    note_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note_paragraph.add_run(note)
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph()


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.04
    paragraph.add_run(text)


def build_document():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    styles["Normal"].font.size = Pt(10.2)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.06

    heading_tokens = [
        ("Heading 1", 15, "2E74B5", 10, 5),
        ("Heading 2", 12.5, "2E74B5", 7, 3),
        ("Heading 3", 11.2, "1F4D78", 4, 2),
    ]
    for name, size, color, before, after in heading_tokens:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "Báo cáo tiến độ reproduce Puppeteer và hướng cải thiện generalization"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "Ngày 20/07/2026"
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Báo cáo tiến độ reproduce Puppeteer và đề xuất cải thiện generalization")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Dựa trên Dang et al. (2025) và khung Combinatorial Generalization của Mahajan et al. (2022)"
    )
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(89, 89, 89)
    subtitle.paragraph_format.space_after = Pt(8)

    callout = doc.add_table(rows=1, cols=1)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, "E8EEF5")
    set_cell_border(cell, color="A9BCD6", sz="8")
    set_cell_margins(cell, top=100, bottom=100, start=150, end=150)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("Tóm tắt ngắn: ")
    run.bold = True
    run.font.color.rgb = RGBColor(31, 58, 95)
    run.font.size = Pt(10)
    run = paragraph.add_run(
        "Đã chạy được pipeline Puppeteer trên GSM-Hard ở quy mô nhỏ, quan sát accuracy trên 38 sample train/test nội bộ có cải thiện, nhưng hiện còn thiếu kiểm tra token cost và các đánh giá generalization khi thay đổi agent pool. Các điểm yếu chính nằm ở agent pool cố định, policy MLP gắn cứng với thứ tự agent, thiếu early stopping, và agent chưa có nhận thức rõ về MAS."
    )
    run.font.size = Pt(10)
    doc.add_paragraph()

    add_paragraph(doc, "1. Tiến độ reproduce code", "Heading 1")
    progress_items = [
        "Đã cấu hình và chạy được pipeline trên GSM-Hard với chế độ initialized/evolved theo checkpoint policy.",
        "Đã train được 38 sample GSM-Hard. Đây là bước xác nhận training loop, logging, checkpoint và embedding-based state representation hoạt động được trong môi trường hiện tại.",
        "Đã test lại trên 38 sample đã train và ghi nhận accuracy tăng so với trước khi train. Token cost chưa được kiểm tra đầy đủ, nên kết luận hiện tại mới dừng ở hiệu quả accuracy nội bộ.",
        "Đã xin đăng ký dùng A100 của trường để mở rộng số sample, giảm thời gian chạy và có điều kiện reproduce bảng kết quả lớn hơn.",
    ]
    for item in progress_items:
        add_bullet(doc, item)

    add_placeholder(
        doc,
        "[Placeholder Hình/Bảng 1] Kết quả reproduce GSM-Hard",
        "Gợi ý chèn: bảng gồm số sample train/test, initialized accuracy, evolved accuracy, token cost trung bình, số agent trung bình/path.",
    )

    add_paragraph(doc, "2. Các vấn đề quan sát được trong implementation hiện tại", "Heading 1")
    issues = [
        (
            "Agent pool cố định và MLP gắn với thứ tự personas",
            "Output neuron thứ i của MLP được map trực tiếp sang agent thứ i trong file personas. Nếu thêm/xóa agent thì output_dim mismatch; nếu đổi thứ tự nhưng giữ số agent, checkpoint có thể vẫn load được nhưng semantic mapping bị sai. Điều này làm policy khó generalize khi thay đổi team composition.",
        ),
        (
            "Không dừng sớm khi non-terminator đã có final answer",
            "Một agent thường có thể sinh FINAL ANSWER, nhưng path chỉ dừng khi policy chọn TerminatorAgent hoặc đạt max_step_num. Điều này gây tốn API/token và có thể làm trajectory dài hơn cần thiết, đặc biệt với GSM-Hard.",
        ),
        (
            "Agent không biết rõ mình đang ở trong MAS",
            "Prompt của từng agent chủ yếu chứa role của chính nó, câu hỏi và kết quả trước đó. Agent gần như không biết agent pool có những role nào, không biết ranh giới nhiệm vụ của mình trong hệ phối hợp, nên dễ làm thay vai trò của agent khác.",
        ),
        (
            "Role specialization chưa sạch",
            "Nhiều prompt yêu cầu cả Planner/Reasoner/Reflect/Summarizer đều kết thúc bằng FINAL ANSWER. Điều này giúp benchmark có đáp án nhưng làm mờ ý nghĩa phối hợp: agent không chỉ tạo intermediate result mà thường tự solve luôn bài toán.",
        ),
    ]
    for title_text, body in issues:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(title_text + ": ")
        run.bold = True
        run.font.color.rgb = RGBColor(31, 77, 120)
        paragraph.add_run(body)

    add_placeholder(
        doc,
        "[Placeholder Hình 2] Sơ đồ điểm yếu hiện tại",
        "Gợi ý chèn: diagram state embedding → MLP output fixed slots → selected agent; highlight lỗi khi đổi personas và khi final answer không trigger stop.",
    )

    add_paragraph(doc, "3. Liên hệ với Mahajan et al. về Combinatorial Generalization", "Heading 1")
    add_paragraph(
        doc,
        "Mahajan et al. xem generalization trong cooperative MAS như khả năng thích nghi khi capability, số lượng agent, hoặc team composition thay đổi. Từ góc nhìn này, Puppeteer của Dang et al. có ưu điểm là điều phối agent động theo state, nhưng vẫn chưa phải một MAS có combinatorial generalization mạnh vì agent pool và action space của policy vẫn cố định.",
    )
    paragraph = doc.add_paragraph()
    paragraph.add_run("Các yêu cầu baseline hiện còn thiếu gồm: ").bold = True
    paragraph.add_run(
        "capability grounding cho từng agent; policy phụ thuộc vào team composition; kiến trúc không gắn cứng với số lượng/thứ tự agent; capability inference từ lịch sử chạy; và benchmark tách rõ seen vs. unseen agent/composition."
    )

    add_paragraph(doc, "4. Đề xuất phương án cải thiện", "Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Hướng cải thiện", "Ý tưởng kỹ thuật", "Kỳ vọng đánh giá"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.2)

    solution_rows = [
        (
            "Capability-aware orchestrator",
            "Thay MLP output cố định bằng scorer f(state, agent_embedding). Agent embedding mã hóa role prompt, model, tools, cost và lịch sử hiệu quả.",
            "Test thêm/xóa/thay agent mà không cần đổi output_dim cố định.",
        ),
        (
            "Permutation-invariant agent encoder",
            "Dùng attention/DeepSets/GNN để aggregate agent pool như một tập, không phụ thuộc thứ tự personas.",
            "So sánh khi shuffle personas, leave-one-agent-out, add-new-agent.",
        ),
        (
            "Early-stop / answer gate",
            "Nếu non-terminator sinh answer hợp lệ với confidence/format đúng, gọi answer verifier hoặc cho phép stop sớm trước max_step_num.",
            "Giảm token cost, số step/path, vẫn giữ hoặc tăng accuracy.",
        ),
        (
            "MAS-aware prompts",
            "Prompt nói rõ agent đang ở trong MAS, danh sách role có sẵn, nhiệm vụ riêng của role, output schema không lấn vai trò agent khác.",
            "Tăng role specialization, giảm duplicated reasoning.",
        ),
    ]
    for row in solution_rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_margins(cells[i])
            run = cells[i].paragraphs[0].add_run(text)
            run.font.size = Pt(8.6)
    set_table_borders(table, color="D9E2F3", sz="6")
    set_col_widths(table, [1.55, 3.05, 1.9])
    doc.add_paragraph()

    add_paragraph(doc, "Phương pháp ưu tiên", "Heading 2")
    add_paragraph(
        doc,
        "Ưu tiên hướng capability-aware orchestrator: thay vì để MLP sinh vector xác suất có kích thước bằng số agent cố định, mỗi agent được biểu diễn bằng một capability embedding. Policy nhận state của task và chấm điểm từng agent bằng cùng một scoring network. Cách này giữ được tinh thần Puppeteer nhưng giảm phụ thuộc vào slot cố định của personas, phù hợp hơn với yêu cầu generalization của Mahajan et al.",
    )

    add_paragraph(doc, "Benchmark đề xuất", "Heading 2")
    benchmark_items = [
        "Seen composition: train/test cùng agent pool để so với baseline Puppeteer hiện tại.",
        "Shuffle personas: giữ agent giống nhau nhưng đổi thứ tự để kiểm tra policy có phụ thuộc index hay không.",
        "Leave-one-agent-out: train thiếu một role/model, test khi thêm lại agent đó.",
        "Add/remove/substitute agent: thêm model mới hoặc thay model của Summarizer/Concluder để đo generalization gap.",
        "Metrics: accuracy, token cost, số step/path, số agent trung bình, early-stop rate, và gap giữa seen vs. unseen composition.",
    ]
    for item in benchmark_items:
        add_bullet(doc, item)

    add_placeholder(
        doc,
        "[Placeholder Bảng 3] Benchmark generalization đề xuất",
        "Gợi ý chèn: các setting Seen / Shuffled / Leave-one-out / Add-agent / Substitute-agent cùng metrics accuracy, token cost, generalization gap.",
    )

    add_paragraph(doc, "Kết luận ngắn", "Heading 1")
    add_paragraph(
        doc,
        "Tiến độ hiện tại đã chứng minh pipeline chạy được và có tín hiệu học trên GSM-Hard ở quy mô nhỏ. Tuy nhiên, để phát triển đóng góp nghiên cứu rõ hơn, bước tiếp theo nên chuyển từ reproduce accuracy sang phân tích generalization khi agent pool thay đổi. Hướng capability-aware, permutation-invariant orchestrator kết hợp early stopping là phương án tự nhiên nhất để nối Dang et al. với Mahajan et al., đồng thời giải quyết trực tiếp các lỗi quan sát được trong code hiện tại.",
    )

    add_paragraph(doc, "Tài liệu tham chiếu", "Heading 2")
    references = [
        "Dang et al. (2025). Multi-Agent Collaboration via Evolving Orchestration.",
        "Mahajan et al. (2022). Generalization in Cooperative Multi-Agent Systems.",
        "Iqbal et al. (2021). Randomized Entity-wise Factorization for Multi-Agent Reinforcement Learning (REFIL).",
        "Ryu et al. (2020). Multi-Agent Actor-Critic with Hierarchical Graph Attention Network.",
    ]
    for reference in references:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.15)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(reference)
        run.font.size = Pt(8.8)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
