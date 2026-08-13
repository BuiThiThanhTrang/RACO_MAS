from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


DOCX_PATH = Path(r"D:\SideProject\ChatDev\output\bao_cao_tien_do_puppeteer_generalization.docx")
IMAGE_PATH = Path(r"D:\SideProject\ChatDev\output\placeholder_2_so_do_diem_yeu_puppeteer.png")


def clear_cell(cell):
    for paragraph in cell.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def insert_image_into_placeholder():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"Missing report DOCX: {DOCX_PATH}")
    if not IMAGE_PATH.exists():
        raise FileNotFoundError(f"Missing placeholder image: {IMAGE_PATH}")

    doc = Document(DOCX_PATH)
    placeholder_text = "[Placeholder Hình 2]"
    replaced = False

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if placeholder_text in cell.text:
                    clear_cell(cell)

                    title = cell.add_paragraph()
                    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = title.add_run("Hình 2. Sơ đồ các điểm yếu trong orchestration hiện tại")
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(31, 77, 120)

                    image_paragraph = cell.add_paragraph()
                    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    image_run = image_paragraph.add_run()
                    image_run.add_picture(str(IMAGE_PATH), width=Inches(6.1))

                    caption = cell.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_run = caption.add_run(
                        "Minh họa pipeline hiện tại: state context được embedding rồi đưa vào MLP fixed slots; "
                        "các callout đỏ đánh dấu điểm yếu cần cải thiện."
                    )
                    caption_run.italic = True
                    caption_run.font.size = Pt(8.5)
                    caption_run.font.color.rgb = RGBColor(89, 89, 89)

                    replaced = True
                    break
            if replaced:
                break
        if replaced:
            break

    if not replaced:
        raise ValueError(f"Could not find placeholder text: {placeholder_text}")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    insert_image_into_placeholder()
