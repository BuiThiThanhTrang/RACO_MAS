from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX_PATH = Path(r"D:\SideProject\ChatDev\output\bao_cao_tien_do_puppeteer_generalization.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2F3", sz="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), sz)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill="F2F4F7", widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.3)
                    if row_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)


def clear_cell(cell):
    for paragraph in list(cell.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    for table in list(cell.tables):
        table._element.getparent().remove(table._element)


def add_caption(cell, text):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(31, 77, 120)
    paragraph.paragraph_format.space_after = Pt(4)


def fill_cell(cell, text, bold=False, color=None):
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(8.3)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_diff_table(cell):
    add_caption(cell, "Bảng 1. Các thay đổi so với thiết lập paper gốc")
    table = cell.add_table(rows=1, cols=4)
    headers = ["Thành phần", "Paper gốc", "Thiết lập hiện tại", "Ảnh hưởng / ghi chú"]
    for idx, header in enumerate(headers):
        fill_cell(table.rows[0].cells[idx], header, bold=True, color="1F4D78")

    rows = [
        [
            "State representation / reward model",
            "Dùng reward model lớn để tạo state representation và reward signal phụ trợ.",
            "Thay bằng embedding API / reward model nhỏ hơn để chạy local và giảm yêu cầu GPU.",
            "Dễ chạy hơn nhưng không hoàn toàn tương đương paper; cần ghi rõ đây là reproduce có điều chỉnh.",
        ],
        [
            "Agent pool và LLM",
            "Agent pool theo paper, gồm các role/model được cấu hình cho Mimas/Titan.",
            "Agent pool đã thay đổi: dùng personas_gsm_local, một số agent/model chuyển sang Gemini/Gemma hoặc bị loại để giảm quota/cost.",
            "Action space và capability distribution khác paper; checkpoint/policy có thể không so sánh trực tiếp.",
        ],
        [
            "Quy mô train",
            "Paper train/evaluate trên tập benchmark lớn hơn.",
            "Hiện mới train được 38 sample GSM-Hard.",
            "Kết quả hiện chỉ là sanity check cho pipeline, chưa đủ kết luận generalization hoặc reproduce đầy đủ.",
        ],
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            fill_cell(cells[idx], text)
    style_table(table, widths=[1.25, 1.65, 1.85, 1.75])


def add_results_table(cell):
    add_caption(cell, "Bảng 2. Kết quả reproduce paper trên GSM-Hard")
    table = cell.add_table(rows=1, cols=7)
    headers = [
        "Setting",
        "Policy mode",
        "#Train samples",
        "#Eval samples",
        "Accuracy",
        "Avg token cost",
        "Ghi chú",
    ]
    for idx, header in enumerate(headers):
        fill_cell(table.rows[0].cells[idx], header, bold=True, color="1F4D78")

    rows = [
        ["Paper - Puppeteer Mono", "Initialized", "—", "—", "0.2467", "—", "Theo bảng paper; cần điền token nếu có."],
        ["Paper - Puppeteer Mono", "Evolved", "—", "—", "0.4800", "—", "Theo bảng paper."],
        ["Paper - Puppeteer", "Initialized", "—", "—", "0.5600", "—", "Theo bảng paper."],
        ["Paper - Puppeteer", "Evolved", "—", "—", "0.5400", "—", "Theo bảng paper."],
        ["Reproduce hiện tại", "Initialized", "0", "38", "[điền]", "[chưa kiểm tra]", "Baseline trước train / cùng 38 sample."],
        ["Reproduce hiện tại", "Evolved/checkpoint", "38", "38", "[điền]", "[chưa kiểm tra]", "Đã quan sát accuracy tăng; cần điền số cụ thể."],
        ["Reproduce mở rộng", "Evolved/checkpoint", "38+", "500 / 1000 / full", "[điền sau]", "[điền sau]", "Sau khi có A100 hoặc quota đủ."],
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            fill_cell(cells[idx], text)
    style_table(table, widths=[1.25, 0.95, 0.75, 0.75, 0.7, 0.85, 1.25])


def replace_placeholder():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(f"Missing report DOCX: {DOCX_PATH}")

    doc = Document(DOCX_PATH)
    placeholder_text = "[Placeholder Hình/Bảng 1]"
    replaced = False

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if placeholder_text in cell.text:
                    clear_cell(cell)
                    add_diff_table(cell)
                    spacer = cell.add_paragraph()
                    spacer.paragraph_format.space_after = Pt(4)
                    add_results_table(cell)
                    replaced = True
                    break
            if replaced:
                break
        if replaced:
            break

    if not replaced:
        raise ValueError(f"Could not find placeholder: {placeholder_text}")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    replace_placeholder()
