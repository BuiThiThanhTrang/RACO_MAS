from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import yaml

from tools.base.base_tool import Tool
from tools.base.register import global_tool_registry


def _load_limits() -> tuple[int, int]:
    config_path = Path(__file__).resolve().parents[1] / "config" / "global.yaml"
    try:
        with config_path.open("r", encoding="utf-8") as source:
            config: dict[str, Any] = yaml.safe_load(source) or {}
        settings = dict(config.get("spreadsheet_tools") or {})
    except (OSError, yaml.YAMLError):
        settings = {}
    return int(settings.get("max_cells", 4000)), int(settings.get("max_chars", 60000))


def _color_value(color) -> str | None:
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    if color_type == "rgb" and getattr(color, "rgb", None):
        return str(color.rgb)
    if color_type == "indexed" and getattr(color, "indexed", None) is not None:
        return f"indexed:{color.indexed}"
    if color_type == "theme" and getattr(color, "theme", None) is not None:
        return f"theme:{color.theme}"
    return None


def _inspect_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("XLSX inspection requires openpyxl") from error

    max_cells, max_chars = _load_limits()
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    value_book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    lines = [f"Workbook: {path.name}", f"Sheets: {', '.join(formula_book.sheetnames)}"]
    emitted = 0

    for sheet_name in formula_book.sheetnames:
        sheet = formula_book[sheet_name]
        value_sheet = value_book[sheet_name]
        lines.extend(
            [
                "",
                f"## Sheet: {sheet_name}",
                f"Used range: A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}",
            ]
        )
        if sheet.merged_cells.ranges:
            lines.append(
                "Merged ranges: " + ", ".join(map(str, sheet.merged_cells.ranges))
            )
        hidden_rows = [str(index) for index, dim in sheet.row_dimensions.items() if dim.hidden]
        hidden_columns = [str(index) for index, dim in sheet.column_dimensions.items() if dim.hidden]
        if hidden_rows:
            lines.append("Hidden rows: " + ", ".join(hidden_rows))
        if hidden_columns:
            lines.append("Hidden columns: " + ", ".join(hidden_columns))

        for row in sheet.iter_rows():
            for cell in row:
                formula = cell.value
                cached = value_sheet[cell.coordinate].value
                fill = _color_value(cell.fill.fgColor) if cell.fill else None
                has_style_evidence = bool(fill and fill not in {"00000000", "FFFFFFFF"})
                if formula is None and not has_style_evidence:
                    continue
                details = [f"{cell.coordinate}={formula!r}"]
                if isinstance(formula, str) and formula.startswith("="):
                    details.append(f"cached={cached!r}")
                if has_style_evidence:
                    details.append(f"fill={fill}")
                if cell.number_format and cell.number_format != "General":
                    details.append(f"number_format={cell.number_format!r}")
                if cell.comment and cell.comment.text:
                    details.append(f"comment={cell.comment.text!r}")
                lines.append(" | ".join(details))
                emitted += 1
                if emitted >= max_cells or sum(map(len, lines)) >= max_chars:
                    lines.append(
                        f"[TRUNCATED after {emitted} populated/styled cells; "
                        "use Python Data Analyst for targeted computation.]"
                    )
                    return "\n".join(lines)[:max_chars]
    return "\n".join(lines)[:max_chars]


def _inspect_csv(path: Path) -> str:
    _, max_chars = _load_limits()
    lines = [f"CSV: {path.name}"]
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        for row_number, row in enumerate(csv.reader(source), start=1):
            lines.append(f"row {row_number}: " + " | ".join(row))
            if sum(map(len, lines)) >= max_chars:
                lines.append(f"[TRUNCATED after row {row_number}]")
                break
    return "\n".join(lines)[:max_chars]


@global_tool_registry("inspect_spreadsheet")
class SpreadsheetInspect(Tool):
    def __init__(self, name: str):
        super().__init__(
            name=name,
            description="inspect CSV/XLSX values, formulas, coordinates, and styles",
            execute_function=self.execute,
        )

    def execute(self, *args, **kwargs):
        file_path = Path(str(kwargs.get("file_path") or "")).resolve()
        if not file_path.is_file():
            return False, f"Spreadsheet file does not exist: {file_path}"
        extension = str(kwargs.get("file_extension") or file_path.suffix).lower()
        try:
            if extension == ".xlsx":
                return True, _inspect_xlsx(file_path)
            if extension == ".csv":
                return True, _inspect_csv(file_path)
            return False, f"Unsupported spreadsheet format: {extension or '[unknown]'}"
        except Exception as error:
            return False, f"Spreadsheet inspection failed ({type(error).__name__}): {error}"
